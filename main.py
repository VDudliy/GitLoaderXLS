import PySimpleGUI as sg
import psycopg2
import openpyxl
import pandas as pd

# Создаем и инициализируем переменные

File_name=''
List_name=''
db_name=''
db_user=''
db_password=''
db_host='localhost'
db_port='5432'
db_table=''
db_type="PostgreSQL"
out_column=''
string_xls=[]
string_db=[]
param_xls=[]
param_db=[]

# Описываем процедуру для Excel
def exel_file():

# Выбираем название файла Excel
    sg.theme('LightGreen')  # please make your creations colorful
    layout = [[sg.Text('Файл XLS ')],
            [sg.Input(key='-file_name-'), sg.FileBrowse()],
            [sg.OK()]]

    window = sg.Window('Выбрать файл XLS для анализа', layout)
    event, values = window.read()
    File_name = values.get('Browse')
    window.close()

#Открываем файл и определяем сколько в нем листов и предлагаем выбрать для дальнейшей работы

    try:
         wb = openpyxl.reader.excel.load_workbook(filename=File_name)
    except:
        return (1)
    sg.theme('LightGreen')
    layout = [[sg.Text('Выберите лист для загрузки')],
          [sg.Listbox(values=wb.sheetnames, size=(30, 12), key='board')],
          [sg.OK('Ok')]]
    window = sg.Window('Выбор листа', layout)
    event, values = window.read()
    mysheet = values.get('board')
    if len(mysheet)==0:
        window.close()
        return (2)
    window.close()

# Определяем столбцы и отбираем для дальнейшей записи в базу

    excel_data_df = pd.read_excel(File_name, sheet_name=mysheet[0])
    mycolumn=excel_data_df.columns.tolist()
    sg.theme('LightGreen')
    layout = [[sg.Text('Выберите столбцы для загрузки')],
          [sg.Listbox(values=mycolumn, size=(30, 12),select_mode="multiple", key='col')],
          [sg.OK('Ok'), sg.Button('Exit')]]
    window = sg.Window('Выбрать cтолбец', layout)
    event, values = window.read()
    mycoll = values.get('col')
    if len(mycoll) == 0:
        window.close()
        return (3)
    mycoll_and_type_out=[]
    for i in mycoll:
        mycoll_and_type = []
        mycoll_and_type.append(i)
        mycoll_and_type.append((excel_data_df[i].dtype.kind))
        mycoll_and_type_out.append(mycoll_and_type)
    window.close()
    return [File_name, mysheet, mycoll_and_type_out]


#Описываем процедуру для подключения к базе данных PostgreSQL

def table_base():
    db_name = ''
    db_user = ''
    db_password = ''
    db_host = 'localhost'
    db_port = '5432'
    db_table = ''
    db_type = "PostgreSQL"
    while True:
          sg.theme('LightGreen')
          layout = [[sg.Text('Выберите тип базы данных и параметры подключения',font='Lucida')],
          [sg.Text('')],
          [sg.Text('Выберите тип базы данных ')],
          [sg.Combo(['PostgreSQL', 'MySql'],
                     default_value='PostgreSQL', key='DB_Type')],
          [sg.Text('Введите наименование базы данных')],
          [sg.Input(db_name, key='db_name')],
          [sg.Text('Введите имя пользователя')],
          [sg.Input(db_user, key='db_user')],
          [sg.Text('Введите пароль пользователя')],
          [sg.Input('', key='db_password')],
          [sg.Text('Введите Host')],
          [sg.Input(db_host, key='db_host')],
          [sg.Text('Введите port')],
          [sg.Input(db_port, key='db_port')],
          [sg.OK('Ok'), sg.Button('Exit')]]
          win = sg.Window('Выбрать базу и соединение', layout)
          e, v = win.read()
          db_name=v.get('db_name')
          db_user=v.get('db_user')
          db_password=v.get('db_password')
          db_host=v.get('db_host')
          db_port=v.get('db_port')
          if e=='Ok':
             win.close()
             try:
                conn = psycopg2.connect(dbname=db_name, user=db_user, password=db_password, host=db_host,port=db_port)
                cursor = conn.cursor()
             except:
                sg.popup_ok('Ошибка подключения к базе данных')
                continue
          else:
              window.close()
              return
          postgres_query = """ SELECT table_name FROM information_schema.tables """
    # where    table_schema = 'public'
          cursor.execute(postgres_query)
          record = cursor.fetchall()

          sg.theme('LightGreen')
          layout = [[sg.Text('Выберите таблицу  базы данных')],
          [sg.Listbox(values=record, size=(30, 12),key='table')],
          [sg.OK('Ok')]]
          window = sg.Window('Выбрать таблицу', layout)
          event, values = window.read()
          db_table=values.get('table')
          window.close()

          postgres_query = """ select column_name,data_type 
          from information_schema.columns 
          where table_name = %s"""

          try:
               cursor.execute(postgres_query,db_table)
          except:
               cursor.close()
               sg.popup_ok('Вы не выбрали таблицу базы данных')
               continue
          record = cursor.fetchall()
          sg.theme('LightGreen')
          layout = [[sg.Text('Выберите поля таблицы')],
          [sg.Listbox(values=record, size=(30, 12),select_mode="multiple",key='coll')],
          [sg.OK('Ok')]]
          window = sg.Window('Выбрать поля', layout)
          event, values = window.read()
          db_call=values.get('coll')
          if len(db_call)==0:
             sg.popup_ok('Не выбраны поля базы данных для загрузки')
             window.close()
             continue
          else:
              window.close()
              return [db_name,db_user,db_password,db_host,db_port,db_table,db_type,db_call]


#Сопоставляем поля XLS и поля Базы данных
#Colunn_base, Column_XLS
def select_column_base_xls(param_xls,param_db):
    sg.theme('LightGreen')  # please make your creations colorful
    param_insert = []
    out_vision = [('', '', '')]
    layout = [
        [sg.Text('Поля XLS для загрузки                                '),
         sg.Text('                          Поля БД для загрузки')],
        [sg.Table(values=param_xls, key='coll_xls',
                  headings=('Имя поля', 'Формат поля'), enable_events=True),
         sg.Table(values=param_db, key='coll_db',
                  headings=('Имя поля', 'Формат поля'), pad=(83, 0), enable_events=True)
         ],
        [sg.Input('', size=(39, 1), key='input_xls'), sg.Button('Добавить связь'),
         sg.Input('', size=(37, 1), key='input_db')],
        [sg.Table(values=out_vision, size=(40, 10), key='key_vision', justification='left', pad=(105, 3),
                  headings=('   Имя поля  XLS  ', '<-->', ' Формат поля  в БД '), enable_events=True),
         sg.Button('Сохранить',size=(10,2),button_color='red')]
    ]

    win3 = sg.Window('Сопоставьте поля Базы данных и Поля  файла XLS ', layout, size=(750, 470))

    out_column = []
    out_prom=[]

    while True:
        event, values = win3.read()
        if event == 'Добавить связь':
            out_prom.append(param_xls[v_xls])
            out_prom.append(list((param_db[v_db])))
            out_column.append(out_prom)
            out_prom=[]
            out_vision.append([v_xls1, '<--->', v_db1])
            param_xls.pop(v_xls)
            param_db.pop(v_db)
            win3['coll_xls'].update(param_xls)
            win3['coll_db'].update(param_db)
            win3['input_xls'].update('')
            win3['input_db'].update('')
            win3['key_vision'].update(out_vision)
            win3['Сохранить'].update(button_color='green')
        elif event == 'coll_xls':
            v_xls = values['coll_xls'][0]
            win3['input_xls'].update(param_xls[v_xls])
            v_xls1 = param_xls[v_xls][0]
        elif event == 'coll_db':
            v_db = values['coll_db'][0]
            win3['input_db'].update(param_db[v_db])
            v_db1 = param_db[v_db][0]
        elif event in (sg.WIN_CLOSED, 'Exit'):
            break
        elif event == ('Сохранить'):
            break

    win3.close()
    return [out_column]



# Запускаем основное меню
sg.theme('KAYAK')
layout = [
    [sg.B('Excel',font='Lucida',size=(15,1)),
    sg.Frame(layout=[
    [sg.Text('Имя файла XLS : '),(sg.Text(size=(25,1), key='-OUTPUT-'))],
    [sg.Text('Имя листа :'),(sg.Text(size=(25,1), key='-OUTPUT1-'))]],
    title='Параметры файла XLS',title_color='GRAY', relief=sg.RELIEF_SUNKEN)],
    [sg.Text(size=(25,1))],
    [sg.B('Database',size=(15,1),font='Lucida'),
    sg.Frame(layout=[
    [sg.Text('Имя базы данных : '), sg.Text(size=(25,1), key='-OUTPUT2-')],
    [sg.Text('Имя таблицы :'), sg.Text(size=(25,1), key='-OUTPUT3-'),]],
    title='Параметры базы данных',title_color='GRAY', relief=sg.RELIEF_SUNKEN)],
    [sg.Text(size=(25,1))],
    [sg.B('Select column',font='Lucida',size=(15,1)),
    sg.Frame(layout=[
    [sg.Text('Создайте связь',key='-OUTPUT4-')],
    [sg.Text('столбцов файла XLS и БД', key='-OUTPUT5-'),]],
    title='Сопоставление столбцов',title_color='GRAY', relief=sg.RELIEF_SUNKEN)],
    [sg.Text(size=(25,1))],
    [sg.Button('Загрузить', size=(100, 2))],
    [sg.Text(size=(25,3))],
    [sg.Button('Exit',size=(15,1))],
    ]
window1 = sg.Window('Выбор файла XLS определение Базы данных для загрузки', layout,size=(470,470))
while True:             # Event Loop
    event, values = window1.read()

    if event == 'Exit':
        window1.close()
        break
    elif event=='Excel':
         window1.hide()
         param_xls=(exel_file())
         if param_xls==1:
             sg.popup_ok("Данные о файле Excel считаны не корректно. Не выбран файл")
         elif param_xls==2:
             sg.popup_ok("Данные о файле Excel считаны не корректно. Не выбран лист")
         elif param_xls==3:
             sg.popup_ok("Данные о файле Excel считаны не корректно. Не выбраны столбцы")
         else:
             window1['-OUTPUT-'].update(param_xls[0])
             window1['-OUTPUT1-'].update(param_xls[1])
         window1.refresh()
         window1.un_hide()

    elif event=='Database':
        window1.hide()
        param_db = (table_base())
        try:
            db_name = param_db[0]
            db_user = param_db[1]
            db_password = param_db[2]
            db_host = param_db[3]
            db_port = param_db[4]
            db_table = param_db[5]
            db_type = param_db[6]
        except:
            sg.popup_ok("Не корректно выбраны параметры базы данных")

        window1['-OUTPUT2-'].update(db_name)
        window1['-OUTPUT3-'].update(db_table)
        window1.refresh()
        window1.un_hide()
    elif event == 'Select column':
         try:
            out_column=select_column_base_xls(param_xls[2],param_db[7])
         except:
             if len(param_xls)==0 or len(param_db)==0:
                 sg.popup_ok('Поля базы данных или файла XLS не определены!!!')
             else:
                if len(param_db[7])==0:
                   sg.popup_ok('Поля базы данных не определены!!!')
                elif len(param_xls[2])==0:
                   sg.popup_ok('Поля файла XLS не определены!!!')
         if out_column!=[]:
            window1['-OUTPUT4-'].update('Связь создана')
            window1['-OUTPUT5-'].update('')
    elif event == 'Загрузить':
        try:
              for i in out_column[0]:
                string_xls.append([i][0][0][0])
                string_db.append([i][0][1][0])
        except:
              sg.popup_ok("Столбцы бызы данных и файла XLS не соотнесены.")
              continue
        excel_data_df = pd.read_excel(param_xls[0], sheet_name=param_xls[1][0])
        Frame_for_write = excel_data_df[string_xls]
        conn = psycopg2.connect(dbname=db_name, user=db_user, password=db_password, host=db_host, port=db_port)
        cur = conn.cursor()
        layout = [[sg.Text('A custom progress meter')],
                  [sg.ProgressBar(len(Frame_for_write), orientation='h', size=(20, 20), key='progressbar')],
                  [sg.Cancel()]]

        # create the window`
        window = sg.Window('Custom Progress Meter', layout)
        progress_bar = window['progressbar']
        for i in range(len(Frame_for_write)):
              Row_for_write=Frame_for_write.iloc[i]
                      # создаем одно место под значение строкив DataFrame
              user_records = ", ".join(["%s"])
                  # данные из Dataframe загоняем в перемеенную
              tuples = []
              tuples.append(tuple(list(Row_for_write)))

                  # создаем список полей в которые будем писать (поля должны соответсвовать полям базы данных или здесь заменить на поля базы данных)
              cols =",".join(string_db)
              query_table=''
              query_table=query_table.join(db_table[0])
                     # создаем запрос
              insert_query = (
                     f"INSERT INTO %s(%s) VALUES %s" %(
              query_table, cols, user_records)
                                  )
              cur.execute(insert_query, tuples)
              conn.commit()
              event, values = window.read(timeout=10)
              if event == 'Cancel' or event == sg.WIN_CLOSED:
                 break
              progress_bar.UpdateBar(i + 1)
        window.close()

    elif event in (sg.WIN_CLOSED, 'Quit'):
        break
window1.close()


